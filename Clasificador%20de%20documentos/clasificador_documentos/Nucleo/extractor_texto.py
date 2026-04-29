"""
Extracción de texto de archivos adjuntos.
Soporta PDF (digital y escaneado via OCR), Word, Excel e imágenes.
Sin dependencias de lógica de clasificación — reutilizable de forma aislada.
"""

from pathlib import Path
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from configuracion.ajustes import TESSERACT_CMD, POPPLER_PATH


def _ocr_pdf(ruta_local: str) -> str:
    """
    Convierte las primeras 3 páginas del PDF a imagen y aplica OCR.
    Requiere: pytesseract, pdf2image, Tesseract instalado y Poppler en PATH.
    Retorna texto vacío si alguna dependencia no está disponible.
    """
    try:
        import pytesseract
        from pdf2image import convert_from_path

        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

        paginas = convert_from_path(
            ruta_local, dpi=200, first_page=1, last_page=3,
            poppler_path=POPPLER_PATH,
        )
        textos = []
        for pagina in paginas:
            texto = pytesseract.image_to_string(pagina, lang="spa+eng+chi_sim")
            textos.append(texto)
        return " ".join(textos)
    except ImportError:
        print("  [OCR] pytesseract o pdf2image no instalados — saltando OCR")
        return ""
    except Exception as e:
        print(f"  [OCR] Error al procesar PDF con OCR: {e}")
        return ""


def _ocr_imagen(ruta_local: str) -> str:
    """
    Aplica OCR directamente sobre una imagen (.jpg, .png, .tiff, .bmp).
    Retorna texto vacío si pytesseract no está disponible.
    """
    try:
        import pytesseract
        from PIL import Image

        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
        img = Image.open(ruta_local)
        return pytesseract.image_to_string(img, lang="spa+eng+chi_sim")
    except ImportError:
        return ""
    except Exception as e:
        print(f"  [OCR] Error al procesar imagen con OCR: {e}")
        return ""


def extraer_texto_por_pagina(ruta_local: str) -> list[dict]:
    """
    Extrae el texto de cada página de un PDF como lista separada.
    Usa pdfplumber; si una página no tiene texto digital aplica OCR página a página.

    Returns:
        Lista de dicts: [{"pagina": 1, "texto": "..."}, ...]
        Lista vacía si no es PDF o si falla la lectura.
    """
    if Path(ruta_local).suffix.lower() != ".pdf":
        return []
    try:
        import pdfplumber
        resultado = []
        with pdfplumber.open(ruta_local) as pdf:
            for i, pagina in enumerate(pdf.pages, start=1):
                texto = pagina.extract_text() or ""
                if not texto.strip():
                    # Página escaneada — OCR individual vía pdf2image
                    try:
                        import pytesseract
                        from pdf2image import convert_from_path
                        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
                        imgs = convert_from_path(
                            ruta_local, dpi=200,
                            first_page=i, last_page=i,
                            poppler_path=POPPLER_PATH,
                        )
                        if imgs:
                            texto = pytesseract.image_to_string(imgs[0], lang="spa+eng")
                    except Exception:
                        pass
                resultado.append({"pagina": i, "texto": texto.strip()})
        return resultado
    except Exception as e:
        print(f"  [WARN] No se pudo leer por pagina {Path(ruta_local).name}: {e}")
        return []


def extraer_texto(ruta_local: str) -> str:
    """
    Extrae texto plano de PDF, Word, Excel o imagen.
    Para PDFs escaneados (sin texto seleccionable), aplica OCR automáticamente.

    Args:
        ruta_local: Ruta al archivo en disco.

    Returns:
        Texto extraído como string. Vacío si no se pudo leer.
    """
    ext = Path(ruta_local).suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(ruta_local) as pdf:
                texto = " ".join(page.extract_text() or "" for page in pdf.pages[:3])
            if texto.strip():
                return texto
            # PDF sin texto seleccionable → escaneado → OCR
            print(f"  [OCR] PDF sin texto digital, aplicando OCR: {Path(ruta_local).name}")
            return _ocr_pdf(ruta_local)

        elif ext == ".docx":
            from docx import Document
            doc = Document(ruta_local)
            return " ".join(p.text for p in doc.paragraphs if p.text.strip())

        elif ext == ".doc":
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            word = None
            try:
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(str(Path(ruta_local).resolve()))
                try:
                    texto = doc.Content.Text
                finally:
                    doc.Close(False)
                return texto
            finally:
                if word is not None:
                    word.Quit()
                pythoncom.CoUninitialize()

        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(ruta_local, read_only=True, data_only=True)
            texto = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    texto.extend(str(c) for c in row if c is not None)
            return " ".join(texto)

        elif ext == ".xls":
            import xlrd
            wb = xlrd.open_workbook(ruta_local)
            texto = []
            for ws in wb.sheets():
                for row_idx in range(ws.nrows):
                    texto.extend(str(c) for c in ws.row_values(row_idx) if c != "")
            return " ".join(texto)

        elif ext in (".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"):
            return _ocr_imagen(ruta_local)

    except Exception as e:
        print(f"  [WARN] No se pudo leer contenido de {Path(ruta_local).name}: {e}")

    return ""
